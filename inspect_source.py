from openpyxl import load_workbook

path = r"D:\Downloads\生产管理进度一览表.xlsx"
wb = load_workbook(path, read_only=True, data_only=False)

for ws in wb.worksheets:
    print(f"SHEET\t{ws.title}\trows={ws.max_row}\tcols={ws.max_column}")
    max_rows = min(ws.max_row or 0, 12)
    max_cols = min(ws.max_column or 0, 30)
    for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
        values = ["" if v is None else str(v).replace("\n", " ")[:80] for v in row]
        if any(values):
            print("ROW\t" + "\t".join(values))
    print()
