from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(r"D:\Downloads\生产管理进度一览表.xlsx")
OUT = Path(r"D:\Documents\New project 2\source_orders.json")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip().replace("\n", " ")


def split_plus(value):
    text = clean(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[+＋]", text) if part.strip()]


def parse_float(text):
    text = clean(text)
    if not text:
        return ""
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else ""


def pick(row, index, *names):
    for name in names:
        if name in index:
            return clean(row[index[name]])
    return ""


def material_type(material, unit):
    s = f"{material} {unit}".upper()
    if "KG" in s or "纱" in material or "丝" in material or "毛" in material:
        return "纱线"
    if "M" in s or "面料" in material or "布" in material or "罗纹" in material or "罗马" in material:
        return "面料"
    return "辅料"


wb = load_workbook(SOURCE, read_only=True, data_only=False)
sheet_name = "2026年度生产进度" if "2026年度生产进度" in wb.sheetnames else wb.sheetnames[-1]
ws = wb[sheet_name]

rows = ws.iter_rows(values_only=True)
headers = [clean(v) for v in next(rows)]
index = {h: i for i, h in enumerate(headers) if h}

orders = []
bom = []
seen_orders = set()

for row in rows:
    order_no = pick(row, index, "合同/销售单号")
    style_no = pick(row, index, "大货款号（品番）", "品番")
    sample_no = pick(row, index, "样品款号")
    if not (order_no or style_no or sample_no):
        continue
    key = order_no or style_no or sample_no
    if key in seen_orders:
        continue
    seen_orders.add(key)

    order = {
        "source_sheet": sheet_name,
        "business": pick(row, index, "业务员", "业务"),
        "register_date": pick(row, index, "登记 日期"),
        "customer": pick(row, index, "客户"),
        "brand": pick(row, index, "品牌"),
        "order_no": order_no or "历史订单",
        "style_no": style_no,
        "sample_no": sample_no,
        "style_name": pick(row, index, "品名（款式）", "款     式"),
        "material_summary": pick(row, index, "素材"),
        "supplier": pick(row, index, "供应商"),
        "order_qty": parse_float(pick(row, index, "订单 数量")),
        "ship_date": pick(row, index, "船期"),
        "purchase_status": pick(row, index, "材料 采购 状态", "大货 意见", "大货意见"),
        "yarn_in_date": pick(row, index, "原料 到库 日期", "原料 到厂 日期", "原料 到厂日"),
        "accessory_in_date": pick(row, index, "辅料 到库 日期", "辅料 到厂日"),
        "factory_knit": pick(row, index, "毛衫 工厂", "生产 工厂"),
        "factory_sew": pick(row, index, "服装 工厂"),
        "factory_finish": pick(row, index, "后道 工厂"),
    }
    orders.append(order)

    mats = split_plus(order["material_summary"]) or [order["material_summary"]]
    qtys = split_plus(pick(row, index, "原料 数量"))
    units = split_plus(pick(row, index, "单位"))
    consumptions = split_plus(pick(row, index, "单耗"))

    max_parts = max(len(mats), len(qtys), len(units), len(consumptions), 1)
    for i in range(max_parts):
        mat = mats[i] if i < len(mats) else (mats[0] if mats else "")
        unit = units[i] if i < len(units) else (units[0] if units else "")
        cons = parse_float(consumptions[i] if i < len(consumptions) else (consumptions[0] if consumptions else ""))
        source_qty = parse_float(qtys[i] if i < len(qtys) else (qtys[0] if qtys else ""))
        if not mat and not unit and not cons:
            continue
        bom.append(
            {
                "bom_id": f"BOM-{len(bom)+1:04d}",
                "order_no": order["order_no"],
                "style_no": style_no,
                "style_name": order["style_name"],
                "material_code": f"MAT-{len(bom)+1:04d}",
                "material_name": mat,
                "material_type": material_type(mat, unit),
                "spec": "",
                "color": "",
                "supplier": order["supplier"],
                "usage_unit": unit or "KG",
                "unit_consumption": cons,
                "order_qty": order["order_qty"],
                "loss_rate": 0.03 if material_type(mat, unit) == "纱线" else 0.05,
                "planned_qty": "",
                "source_required_qty": source_qty,
            }
        )

    if len(orders) >= 80:
        break

OUT.write_text(json.dumps({"orders": orders, "bom": bom}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"wrote {OUT} orders={len(orders)} bom={len(bom)}")
